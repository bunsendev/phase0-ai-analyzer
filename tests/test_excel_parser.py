from pathlib import Path

from openpyxl import Workbook

from phase0_analyzer.parsers.excel import ExcelParser


def test_excel_parser_reads_multiple_sheets_headers_formulas_and_merges(
    tmp_path: Path,
) -> None:
    path = tmp_path / "orders.xlsx"
    workbook = Workbook()
    orders = workbook.active
    orders.title = "受注"
    orders.append(["注文番号", "金額", "税込"])
    orders.append(["A-001", 100, "=B2*1.1"])
    orders.append([None, None, None])
    orders.merge_cells("A4:B4")
    orders["A4"] = "合計"
    inventory = workbook.create_sheet("在庫")
    inventory.append(["商品", "数量"])
    inventory.append(["商品A", 5])
    workbook.save(path)
    workbook.close()

    result = ExcelParser(preview_max_rows=30).parse(11, path.name, path)

    assert result.success
    assert result.file_id == 11
    assert result.metadata["sheet_names"] == ["受注", "在庫"]
    sheets = {sheet["sheet_name"]: sheet for sheet in result.metadata["sheets"]}
    assert sheets["受注"]["header_candidate"] == ["注文番号", "金額", "税込"]
    assert sheets["受注"]["row_count"] == 4
    assert sheets["受注"]["column_count"] == 3
    assert sheets["受注"]["has_formulas"] is True
    assert sheets["受注"]["has_merged_cells"] is True
    assert sheets["受注"]["empty_row_count"] == 1
    assert len(result.table_preview) <= 30


def test_broken_xlsx_returns_failure_result(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not an xlsx file")

    result = ExcelParser().parse(1, path.name, path)

    assert not result.success
    assert result.error_code == "EXCEL_READ_FAILED"
