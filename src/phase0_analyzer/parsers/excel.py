"""Excel workbook parser."""

from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from phase0_analyzer.parsers.base import BaseParser
from phase0_analyzer.parsers.models import ParseError, ParseResult


MAX_SOURCE_ROWS = 100


def _serializable_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


class ExcelParser(BaseParser):
    """Extract workbook structure and representative cell values."""

    parser_name = "excel"
    supported_extensions = frozenset({"xlsx"})

    def __init__(self, preview_max_rows: int = 30) -> None:
        self.preview_max_rows = preview_max_rows

    def _parse(self, file_id: int, file_name: str, path: Path) -> ParseResult:
        try:
            workbook = load_workbook(path, read_only=False, data_only=False)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
            raise ParseError(
                "EXCEL_READ_FAILED", "Excelファイルを読み取れませんでした。"
            ) from error

        sheet_metadata: list[dict[str, Any]] = []
        table_preview: list[dict[str, Any]] = []
        extracted_parts: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                max_row = worksheet.max_row
                max_column = worksheet.max_column
                rows: list[list[Any]] = []
                header_candidate: list[Any] = []
                empty_row_count = 0

                for row_number in range(1, max_row + 1):
                    values = [
                        _serializable_value(
                            worksheet.cell(row=row_number, column=column_number).value
                        )
                        for column_number in range(1, max_column + 1)
                    ]
                    if not any(value not in (None, "") for value in values):
                        empty_row_count += 1
                    elif not header_candidate:
                        header_candidate = values

                    if row_number <= MAX_SOURCE_ROWS:
                        rows.append(values)
                        extracted_parts.extend(
                            str(value) for value in values if value not in (None, "")
                        )
                        if len(table_preview) < self.preview_max_rows:
                            table_preview.append(
                                {
                                    "sheet_name": worksheet.title,
                                    "row_number": row_number,
                                    "values": values,
                                }
                            )

                empty_column_count = sum(
                    all(
                        worksheet.cell(row=row_number, column=column_number).value
                        in (None, "")
                        for row_number in range(1, max_row + 1)
                    )
                    for column_number in range(1, max_column + 1)
                )
                has_formulas = any(
                    worksheet.cell(row=row_number, column=column_number).data_type == "f"
                    for row_number in range(1, max_row + 1)
                    for column_number in range(1, max_column + 1)
                )
                sheet_metadata.append(
                    {
                        "sheet_name": worksheet.title,
                        "row_count": max_row,
                        "column_count": max_column,
                        "header_candidate": header_candidate,
                        "first_100_rows": rows,
                        "has_formulas": has_formulas,
                        "has_merged_cells": bool(worksheet.merged_cells.ranges),
                        "empty_row_count": empty_row_count,
                        "empty_column_count": empty_column_count,
                    }
                )
        finally:
            workbook.close()

        return ParseResult(
            file_id=file_id,
            file_name=file_name,
            file_type="xlsx",
            parser_name=self.parser_name,
            success=True,
            extracted_text="\n".join(extracted_parts),
            table_preview=table_preview,
            metadata={
                "sheet_names": [sheet["sheet_name"] for sheet in sheet_metadata],
                "sheets": sheet_metadata,
            },
        )
